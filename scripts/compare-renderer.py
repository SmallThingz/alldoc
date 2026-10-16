#!/usr/bin/env python3
"""Isolated host s-vs3 renderer builds; timing requires an external quiet/guarded run.

Build does not change packaged jniLibs. Fixture tools are developer-only Python
packages: python-docx, python-pptx, openpyxl and Pillow. Run from the repo root.
"""
import argparse
import datetime
import fcntl
import hashlib
import importlib.metadata
import importlib.util
import io
import json
import os
from pathlib import Path
import resource
import signal
import shutil
import statistics
import struct
import subprocess
import sys
import time
import zipfile

BASE = Path('.tmp/performance/renderer')
MANIFEST = Path('native/render/core/Cargo.toml')
KINDS = {'docx': 1, 'pptx': 2, 'xlsx': 3, 'psd': 4, 'tiff': 5}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2) + '\n')


def native_builder():
    spec = importlib.util.spec_from_file_location('alldoc_native_build', 'native/render/build.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def fixtures():
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from PIL import Image
    directory = BASE / 'fixtures'
    directory.mkdir(parents=True, exist_ok=True)
    for path in Path('app/src/androidTest/assets/office').iterdir():
        if path.suffix[1:] in KINDS:
            shutil.copyfile(path, directory / path.name)
    stamp = datetime.datetime(2026, 1, 1, tzinfo=datetime.timezone.utc)
    doc = Document()
    doc.core_properties.created = stamp
    doc.core_properties.modified = stamp
    doc.add_heading('Field notes and measurements', 0)
    for i in range(100):
        if i % 20 == 0:
            doc.add_heading(f'Section {i // 20 + 1}', 1)
        p = doc.add_paragraph()
        p.add_run(f'Observation {i + 1}. ').bold = True
        p.add_run(('Café field notes compare dimensions, material and conditions. '
                   'A second reading checks the result and records remaining questions. ') * 3)
    table = doc.add_table(rows=21, cols=4)
    for row, cells in enumerate(table.rows):
        for column, cell in enumerate(cells.cells):
            cell.text = f'{row * 4 + column}: measured value'
    doc.save(directory / 'report.docx')
    deck = Presentation()
    deck.core_properties.created = stamp
    deck.core_properties.modified = stamp
    for i in range(15):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f'Field observations {i + 1}'
        slide.placeholders[1].text = '\n'.join(f'Sample {j + 1}: dimensions, material and measured conditions.' for j in range(7))
        box = slide.shapes.add_textbox(Inches(0.6), Inches(6.7), Inches(8), Inches(0.4))
        box.text = f'Progress review — section {i + 1} of 15'
    deck.save(directory / 'review.pptx')
    book = Workbook()
    book.properties.created = stamp
    book.properties.modified = stamp
    sheet = book.active
    sheet.title = 'Measurements'
    sheet.append(['Sample'] + [f'Reading {i}' for i in range(1, 12)])
    for row in range(240):
        sheet.append([f'Sample {row + 1}'] + [round((row + 1) * (column + 2) / 7, 2) for column in range(11)])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='264653')
    sheet.print_title_rows = '1:1'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    book.save(directory / 'measurements.xlsx')
    image = Image.new('RGB', (1024, 768))
    image.putdata([(x % 256, y % 256, (x + y) % 256) for y in range(768) for x in range(1024)])
    image.save(directory / 'scan.tiff', compression='tiff_lzw')
    width, height = 640, 480
    header = b'8BPS' + struct.pack('>H6sHIIHH', 1, b'\0' * 6, 3, height, width, 8, 3)
    planes = b''.join(bytes((x + channel * 47) % 256 for y in range(height) for x in range(width)) for channel in range(3))
    (directory / 'composite.psd').write_bytes(header + struct.pack('>IIIH', 0, 0, 0, 0) + planes)
    # Freeze ZIP entry times too; byte hashes guard the exact corpus during timing.
    for path in directory.iterdir():
        if path.suffix[1:] not in ('docx', 'pptx', 'xlsx'):
            continue
        with zipfile.ZipFile(path) as archive:
            entries = [(item.filename, archive.read(item)) for item in archive.infolist()]
        with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries:
                item = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
                item.compress_type = zipfile.ZIP_DEFLATED
                archive.writestr(item, data)
    return {str(path): digest(path) for path in sorted(directory.iterdir())}


def build(args):
    BASE.mkdir(parents=True, exist_ok=True)
    builder = native_builder()
    builder.prepare()
    sources = builder.source_hashes()
    corpus = fixtures()
    env = os.environ.copy()
    for key in list(env):
        if key in ('RUSTFLAGS', 'CARGO_ENCODED_RUSTFLAGS', 'CRAFT_FONTS_DIR', 'CARGO_TARGET_DIR') or key.startswith('CARGO_PROFILE_RELEASE_'):
            env.pop(key)
    env['CARGO_BUILD_JOBS'] = str(args.jobs)
    compiler = subprocess.check_output(['rustc', '-Vv'], text=True)
    binaries = {}
    for profile in ('s', '3'):
        target = BASE / ('profile-' + profile)
        env['CARGO_PROFILE_RELEASE_OPT_LEVEL'] = profile
        command = ['cargo', 'build', '--locked', '--release', '--bin', 'render', '--manifest-path', str(MANIFEST), '--target-dir', str(target)]
        with (BASE / ('build-' + profile + '.log')).open('w') as log:
            subprocess.run(command, env=env, stdout=log, stderr=subprocess.STDOUT, check=True)
        path = target / 'release/render'
        binaries[profile] = {'path': str(path), 'sha256': digest(path), 'bytes': path.stat().st_size}
        if builder.source_hashes() != sources:
            raise RuntimeError('Native sources changed during comparison build')
    launcher = BASE / 'allrun'
    subprocess.run(['cc', '-std=c11', '-O2', '-DNDEBUG', '-fPIE', '-pie', '-fstack-protector-strong', '-D_FORTIFY_SOURCE=2',
                    '-ffunction-sections', '-fdata-sections', '-Wall', '-Wextra', '-Werror',
                    '-Wl,--gc-sections,-z,relro,-z,now,-z,noexecstack,-z,max-page-size=16384',
                    'native/runner/allrun.c', '-o', str(launcher)], check=True)
    receipt = {'sources': sources, 'fixtures': corpus, 'binaries': binaries, 'launcher': {'path': str(launcher), 'sha256': digest(launcher)},
               'rustc': compiler, 'target': 'host; Android confirmation required', 'lto': 'thin', 'codegenUnits': 1, 'panic': 'unwind',
               'fixtureTools': {name: importlib.metadata.version(name) for name in ('python-docx', 'python-pptx', 'openpyxl', 'Pillow')},
               'pins': {name: value['commit'] for name, value in json.loads(Path('native/render/projection.json').read_text()).items()}}
    save(BASE / 'build.json', receipt)
    print(json.dumps({'built': binaries, 'receipt': str(BASE / 'build.json')}), flush=True)


def verify(receipt):
    builder = native_builder()
    builder.prepare()
    if builder.source_hashes() != receipt['sources']:
        raise RuntimeError('Native sources differ from the comparison build')
    for path, expected in receipt['fixtures'].items():
        if digest(Path(path)) != expected:
            raise RuntimeError('Benchmark fixture changed: ' + path)
    for item in [*receipt['binaries'].values(), receipt['launcher']]:
        if digest(Path(item['path'])) != item['sha256']:
            raise RuntimeError('Benchmark binary changed: ' + item['path'])


def decode(path):
    from PIL import Image
    data = path.read_bytes()
    if len(data) < 32 or len(data) > 16 * 1024 * 1024 + 8224 or data[:4] != b'ZFR1':
        raise RuntimeError('Invalid renderer receipt')
    kind, width, height, page, pages, note, size = struct.unpack_from('<7I', data, 4)
    if kind != 1 or not (1 <= width <= 1400 and 1 <= height <= 1400 and 0 <= page < pages <= 128) or 32 + note + size != len(data):
        raise RuntimeError('Renderer rejected workload or returned invalid metadata: ' + data[32:32 + min(note, 512)].decode(errors='replace'))
    with Image.open(io.BytesIO(data[32 + note:])) as image:
        if image.size != (width, height):
            raise RuntimeError('PNG dimensions differ from receipt')
        pixels = hashlib.sha256(image.convert('RGBA').tobytes()).hexdigest()
    return {'width': width, 'height': height, 'page': page, 'pages': pages, 'pixels': pixels}


def measure(args):
    run_id = datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%S.%fZ') + '-' + str(os.getpid())
    final = BASE / 'timings.json'
    partial = BASE / 'timings.partial.json'
    BASE.mkdir(parents=True, exist_ok=True)
    if final.exists():
        final.replace(BASE / 'timings.previous.json')
    results = []
    save(partial, {'runId': run_id, 'status': 'running', 'results': results})
    try:
        measure_run(args, run_id, results)
    except BaseException as failure:
        final.unlink(missing_ok=True)
        save(partial, {'runId': run_id, 'status': 'failed', 'error': f'{type(failure).__name__}: {failure}', 'results': results})
        raise
    partial.unlink(missing_ok=True)


def measure_run(args, run_id, results):
    if not args.guarded:
        raise RuntimeError('Timing requires host-quiet, then guarded-run --abort-on-busy around this command with --guarded')
    receipt = json.loads((BASE / 'build.json').read_text())
    verify(receipt)
    env = os.environ.copy()
    env.pop('CRAFT_FONTS_DIR', None)
    env.update(LC_ALL='C.UTF-8', RAYON_NUM_THREADS='2')
    run_index = 0
    def expired(*_):
        raise TimeoutError('Renderer exceeded the 16-second benchmark deadline')
    def run(profile, source, side, page):
        nonlocal run_index
        run_index += 1
        output = BASE / ('receipt-' + str(run_index))
        command = [str(Path(receipt['launcher']['path']).resolve()), str(os.getpid()),
                   str(Path(receipt['binaries'][profile]['path']).resolve()), str(KINDS[source.suffix[1:]]),
                   str(source.resolve()), str(output.resolve()), str(side), str(page)]
        previous = signal.signal(signal.SIGALRM, expired)
        signal.setitimer(signal.ITIMER_REAL, 16)
        before = resource.getrusage(resource.RUSAGE_CHILDREN)
        start = time.perf_counter_ns()
        try:
            # Blocking wait avoids Python timeout polling quantizing short timings.
            result = subprocess.run(command, env=env, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            wall = time.perf_counter_ns() - start
            after = resource.getrusage(resource.RUSAGE_CHILDREN)
            if result.returncode != 0:
                raise RuntimeError(f'{source.name}, profile {profile}: renderer exit {result.returncode}')
            decoded = decode(output)
            if decoded['page'] != page or max(decoded['width'], decoded['height']) > side:
                raise RuntimeError('Renderer returned the wrong page or output size')
            return {'wallNs': wall, 'cpuNs': round(((after.ru_utime + after.ru_stime) - (before.ru_utime + before.ru_stime)) * 1e9), 'image': decoded}
        finally:
            signal.setitimer(signal.ITIMER_REAL, 0)
            signal.signal(signal.SIGALRM, previous)
            output.unlink(missing_ok=True)
    for filename in receipt['fixtures']:
        source = Path(filename)
        for side in (1400, 128):
            first = {profile: run(profile, source, side, 0) for profile in ('s', '3')}
            if first['s']['image'] != first['3']['image']:
                raise RuntimeError('First-page profile output mismatch: ' + filename)
            pages = [0, 1] if first['s']['image']['pages'] > 1 else [0]
            for page in pages:
                reference = first['s']['image'] if page == 0 else run('s', source, side, page)['image']
                trials = {'s': [], '3': []}
                for trial in range(args.trials):
                    for profile in (('s', '3') if trial % 2 == 0 else ('3', 's')):
                        measured = run(profile, source, side, page)
                        if measured['image'] != reference:
                            raise RuntimeError(f'Pixel/page mismatch: {filename}, page {page}, profile {profile}')
                        trials[profile].append({key: measured[key] for key in ('wallNs', 'cpuNs')})
                row = {'fixture': filename, 'side': side, 'page': page, 'firstProcess': first if page == 0 else None,
                       'trials': trials, 'median': {profile: {key: statistics.median(t[key] for t in values) for key in ('wallNs', 'cpuNs')}
                                                   for profile, values in trials.items()}}
                results.append(row)
                save(BASE / 'timings.partial.json', {'runId': run_id, 'status': 'running', 'results': results})
    verify(receipt)
    save(BASE / 'timings.json', {'runId': run_id, 'status': 'complete', 'build': receipt, 'externalGuardRequired': True, 'trials': args.trials, 'results': results,
                                'limits': 'Host process/render timings, not disk-cold or UI/cache/Android performance; contamination invalidates this file.'})
    print(str(BASE / 'timings.json'))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    make = commands.add_parser('build', help='Build only; no benchmark or packaged library changes')
    make.add_argument('--jobs', type=int, choices=range(1, 5), default=2)
    timed = commands.add_parser('measure', help='Root must wrap this command with host-quiet/guarded-run')
    timed.add_argument('--guarded', action='store_true')
    timed.add_argument('--trials', type=int, choices=range(7, 31), default=7)
    args = parser.parse_args()
    if not sys.platform.startswith('linux') or not MANIFEST.is_file():
        parser.error('Run on the Linux build host from the AllDoc repository root')
    lock = Path('.tmp/native/render/build.lock')
    lock.parent.mkdir(parents=True, exist_ok=True)
    with lock.open('w') as held:
        fcntl.flock(held, fcntl.LOCK_EX | fcntl.LOCK_NB)
        build(args) if args.command == 'build' else measure(args)


if __name__ == '__main__':
    main()
