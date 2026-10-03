#!/usr/bin/env python3
"""Small synthetic interoperability corpus. Requires python-docx, openpyxl, python-pptx, reportlab."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from reportlab.pdfgen import canvas
root=Path(__file__).resolve().parents[1]/'fixtures';root.mkdir(exist_ok=True)
d=Document();sec=d.sections[0];sec.header.paragraphs[0].text='ALLDOC / FIELD NOTES'
d.add_heading('A little room to think.',0)
p=d.add_paragraph('Project notebook  /  October 2026');p.runs[0].font.color.rgb=RGBColor.from_string('235A49')
d.add_paragraph('A document should feel as clear as the idea inside it. This synthetic file tests paragraphs, emphasis, tables, headers and page breaks.')
p=d.add_paragraph();p.add_run('Keep the original. ').bold=True;p.add_run('Make a copy. ').italic=True;p.add_run('Review the result.')
d.add_heading('What matters',1)
for x in ['Quiet, readable tools','Predictable saving','Honest format boundaries']:d.add_paragraph(x,style='List Bullet')
t=d.add_table(rows=1, cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Workstream','Owner','Status']):c.text=v
for values in [['Canvas','Design','In review'],['Roundtrip','Engine','Testing'],['Offline','Platform','Ready']]:
 for c,v in zip(t.add_row().cells,values):c.text=v
d.add_heading('Across languages',1);d.add_paragraph('English · العربية · 日本語 · हिन्दी · Español')
d.add_page_break();d.add_heading('The next page',1);d.add_paragraph('Explicit page breaks should remain visible. Layout varies with available fonts; the original file remains unchanged.')
sec.footer.paragraphs[0].text='Synthetic test document • No private information'
d.save(root/'field-notes.docx')
w=Workbook();s=w.active;s.title='Project budget';s.append(['Item','Quantity','Cost','Total']);s.append(['Paper',3,4.50,'=B2*C2']);s.append(['Pens',12,1.75,'=B3*C3']);s['A8']='مرحبا';s['B8']='世界';w.create_sheet('Archive').append(['Old entry','Preserved']);w.save(root/'budget.xlsx')
c=canvas.Canvas(str(root/'reading.pdf'));c.setFont('Helvetica-Bold',28);c.drawString(64,760,'A little room to think.');c.setFont('Helvetica',12);c.drawString(64,725,'Synthetic PDF for page navigation and zoom tests.');c.showPage();c.drawString(64,760,'Second page');c.save()
(root/'notes.txt').write_text('A little room to think.\n\nEnglish · العربية · 日本語\n',encoding='utf-8')

# Presentation order intentionally differs from the numbered slide part names.
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt as SlidePt
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED, ZIP_STORED
from xml.sax.saxutils import escape
prs=Presentation();prs.slide_width=SlideInches(13.333333);prs.slide_height=SlideInches(7.5)
for title,body in [('A sharper focus.','Review / October 2026\nA compact place for documents, sheets and ideas.'),('Small by design.','Offline first\nKeep the original\nExport a considered copy'),('Work that moves.','01  Read with context\n02  Make focused changes\n03  Share a fresh copy')]:
 slide=prs.slides.add_slide(prs.slide_layouts[6])
 title_box=slide.shapes.add_textbox(SlideInches(.8),SlideInches(.8),SlideInches(11.6),SlideInches(1.2))
 title_box.text_frame.text=title;title_box.text_frame.paragraphs[0].font.size=SlidePt(36)
 body_box=slide.shapes.add_textbox(SlideInches(.85),SlideInches(2.6),SlideInches(10.8),SlideInches(3.0))
 body_box.text_frame.text=body
 for paragraph in body_box.text_frame.paragraphs:paragraph.font.size=SlidePt(22)
ids=prs.slides._sldIdLst;ids.insert(1,ids[2])
prs.save(root/'studio-review.pptx')

ODF_NAMESPACES="""xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
 xmlns:of="urn:oasis:names:tc:opendocument:xmlns:of:1.2"
"""
def odf_content(body):
 return '<?xml version="1.0" encoding="UTF-8"?><office:document-content '+ODF_NAMESPACES+' office:version="1.3"><office:body>'+body+'</office:body></office:document-content>'
def odf_package(name,kind,content,styles=None):
 mime='application/vnd.oasis.opendocument.'+kind
 with ZipFile(root/name,'w',ZIP_DEFLATED) as out:
  info=ZipInfo('mimetype');info.compress_type=ZIP_STORED;out.writestr(info,mime)
  out.writestr('content.xml',content)
  members=['content.xml']
  if styles:out.writestr('styles.xml',styles);members.append('styles.xml')
  out.writestr('META-INF/manifest.xml','<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3"><manifest:file-entry manifest:full-path="/" manifest:media-type="'+mime+'"/>'+''.join('<manifest:file-entry manifest:full-path="'+part+'" manifest:media-type="text/xml"/>' for part in members)+'</manifest:manifest>')
notes=odf_content('<office:text><text:h text:outline-level="1">Make room for the idea.</text:h><text:p>A clear page, a thoughtful edit, and an original kept safe.</text:p><text:h text:outline-level="2">Working notes</text:h><text:p>Read first. Export a copy. Review the result.</text:p><text:p>English · العربية · 日本語 · Español</text:p></office:text>')
odf_package('studio-notes.odt','text',notes)
(root/'flat-notes.fodt').write_text(notes.replace('office:document-content','office:document'),encoding='utf-8')
def cell(text):return '<table:table-cell office:value-type="string"><text:p>'+escape(text)+'</text:p></table:table-cell>'
sheet='<office:spreadsheet><table:table table:name="Studio budget"><table:table-row>'+''.join(cell(s) for s in ['Item','Quantity','Cost','Total'])+'</table:table-row>'
for name,quantity,cost in [('Notebooks',6,'8.5'),('Pens',12,'2.25')]:
 row=2 if name=='Notebooks' else 3
 sheet+='<table:table-row>'+cell(name)+'<table:table-cell office:value-type="float" office:value="'+str(quantity)+'"/><table:table-cell office:value-type="float" office:value="'+cost+'"/><table:table-cell table:formula="of:=[.B'+str(row)+']*[.C'+str(row)+']" office:value-type="float" office:value="'+str(quantity*float(cost))+'"/></table:table-row>'
sheet+='<table:table-row table:number-rows-repeated="3"><table:table-cell table:number-columns-repeated="4"/></table:table-row><table:table-row>'+cell('Budget notes')+cell('Original retained')+'</table:table-row></table:table><table:table table:name="Archive"><table:table-row>'+cell('Old entry')+cell('Preserved')+'</table:table-row></table:table></office:spreadsheet>'
odf_package('studio-budget.ods','spreadsheet',odf_content(sheet))
slides='<office:presentation>'
for i,(title,body) in enumerate([('A sharper focus.','A compact workspace for documents and ideas.'),('Work that moves.','Read. Refine. Export a copy.')],1):
 slides+='<draw:page draw:name="Slide '+str(i)+'" draw:master-page-name="Default"><draw:frame svg:x="0.8in" svg:y="0.8in" svg:width="11.5in" svg:height="1.1in"><draw:text-box><text:p>'+title+'</text:p></draw:text-box></draw:frame><draw:frame svg:x="0.8in" svg:y="2.5in" svg:width="10.8in" svg:height="3in"><draw:text-box><text:p>'+body+'</text:p></draw:text-box></draw:frame></draw:page>'
slides+='</office:presentation>'
styles='<office:document-styles '+ODF_NAMESPACES+' office:version="1.3"><office:automatic-styles><style:page-layout style:name="wide"><style:page-layout-properties fo:page-width="13.333333in" fo:page-height="7.5in"/></style:page-layout></office:automatic-styles><office:master-styles><style:master-page style:name="Default" style:page-layout-name="wide"/></office:master-styles></office:document-styles>'
odf_package('studio-review.odp','presentation',odf_content(slides),styles)
(root/'manifest.xml').write_text('<?xml version="1.0" encoding="UTF-8"?>\n<workspace name="alldoc">\n  <document format="ODF">Studio notes</document>\n  <policy>Keep the original</policy>\n</workspace>\n',encoding='utf-8')
